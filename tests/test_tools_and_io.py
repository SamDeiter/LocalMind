"""
Comprehensive tests for LocalMind document tools, atomic IO, and error strategy.

Covers:
  1. backend.core.atomic_io   — AtomicFileWriter, validation, SHA-256
  2. backend.core.error_strategy — ErrorCategory, classify_error, DegradationCascade
  3. backend.tools.pptx_tool  — PowerPoint read/edit actions, path jailing
  4. backend.tools.excel_tool — Excel read/write actions, path safety
  5. backend.tools.word_tool  — Word document actions, path safety
  6. backend.tools.pdf_tool   — PDF extraction actions, path safety
"""

import asyncio
import hashlib
import json
import sys
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

# Ensure project root is on sys.path
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


# ═══════════════════════════════════════════════════════════════════════════════
# 1. AtomicFileWriter tests
# ═══════════════════════════════════════════════════════════════════════════════

from backend.core.atomic_io import (
    AtomicFileWriter,
    FileWriteResult,
    _validate_docx,
    _validate_pptx,
    _validate_xlsx,
    compute_sha256,
    validate_file,
)


class TestAtomicFileWriterBasic:
    """Core write / result contract."""

    def test_write_atomic_creates_file(self, tmp_path):
        """Successful atomic write produces the target file."""
        target = tmp_path / "output.txt"
        writer = AtomicFileWriter()

        result = writer.write_atomic(
            target, lambda tmp: tmp.write_text("hello world")
        )

        assert target.exists()
        assert target.read_text() == "hello world"
        assert isinstance(result, FileWriteResult)
        assert result.path == target

    def test_write_atomic_returns_correct_size(self, tmp_path):
        """size_bytes in FileWriteResult matches the actual file size."""
        target = tmp_path / "sized.bin"
        data = b"x" * 1024
        writer = AtomicFileWriter()

        result = writer.write_atomic(
            target, lambda tmp: tmp.write_bytes(data)
        )

        assert result.size_bytes == 1024
        assert target.stat().st_size == 1024

    def test_write_atomic_sha256_matches(self, tmp_path):
        """SHA-256 in FileWriteResult matches an independent hash of the file."""
        target = tmp_path / "hashed.txt"
        content = b"deterministic content for hashing"
        writer = AtomicFileWriter()

        result = writer.write_atomic(
            target, lambda tmp: tmp.write_bytes(content)
        )

        expected = hashlib.sha256(content).hexdigest()
        assert result.sha256 == expected

    def test_write_atomic_creates_parent_dirs(self, tmp_path):
        """AtomicFileWriter creates parent directories if they do not exist."""
        target = tmp_path / "deep" / "nested" / "dir" / "output.txt"
        writer = AtomicFileWriter()

        writer.write_atomic(
            target, lambda tmp: tmp.write_text("nested")
        )

        assert target.exists()
        assert target.read_text() == "nested"

    def test_write_atomic_overwrites_existing(self, tmp_path):
        """Atomic write replaces existing file content."""
        target = tmp_path / "overwrite.txt"
        target.write_text("old content")
        writer = AtomicFileWriter()

        writer.write_atomic(
            target, lambda tmp: tmp.write_text("new content")
        )

        assert target.read_text() == "new content"

    def test_write_atomic_cleans_up_on_write_fn_failure(self, tmp_path):
        """If write_fn raises, the original file is untouched and temp is cleaned."""
        target = tmp_path / "stable.txt"
        target.write_text("original")
        writer = AtomicFileWriter()

        with pytest.raises(RuntimeError, match="boom"):
            writer.write_atomic(
                target, lambda tmp: (_ for _ in ()).throw(RuntimeError("boom"))
            )

        assert target.read_text() == "original"
        # No temp files left behind
        siblings = list(tmp_path.glob("stable.tmp.*"))
        assert siblings == []


class TestAtomicFileWriterValidation:
    """Format validation and rollback on invalid output."""

    def test_json_validation_passes(self, tmp_path):
        """Valid JSON passes validation and write succeeds."""
        target = tmp_path / "good.json"
        writer = AtomicFileWriter()

        result = writer.write_atomic(
            target, lambda tmp: tmp.write_text('{"key": "value"}')
        )

        assert result.path == target
        assert json.loads(target.read_text()) == {"key": "value"}

    def test_json_validation_fails_rolls_back(self, tmp_path):
        """Invalid JSON triggers ValueError and original file is preserved."""
        target = tmp_path / "bad.json"
        target.write_text('{"original": true}')
        writer = AtomicFileWriter()

        with pytest.raises(ValueError, match="not valid JSON"):
            writer.write_atomic(
                target, lambda tmp: tmp.write_text("NOT JSON {{{")
            )

        assert json.loads(target.read_text()) == {"original": True}

    def test_empty_file_fails_validation(self, tmp_path):
        """A zero-byte write is rejected by validate_file."""
        target = tmp_path / "empty.txt"
        writer = AtomicFileWriter()

        with pytest.raises(ValueError, match="empty"):
            writer.write_atomic(target, lambda tmp: tmp.write_bytes(b""))

    def test_pptx_validation_rejects_non_zip(self, tmp_path):
        """_validate_pptx rejects a file that is not a ZIP archive."""
        bad = tmp_path / "bad.pptx"
        bad.write_text("not a zip")

        with pytest.raises(ValueError, match="not a valid ZIP"):
            _validate_pptx(bad)

    def test_pptx_validation_rejects_missing_content_types(self, tmp_path):
        """_validate_pptx rejects a ZIP that lacks [Content_Types].xml."""
        bad = tmp_path / "missing_ct.pptx"
        with zipfile.ZipFile(bad, "w") as zf:
            zf.writestr("random.xml", "<x/>")

        with pytest.raises(ValueError, match="Content_Types"):
            _validate_pptx(bad)

    def test_xlsx_validation_rejects_missing_workbook(self, tmp_path):
        """_validate_xlsx rejects a ZIP missing xl/workbook.xml."""
        bad = tmp_path / "bad.xlsx"
        with zipfile.ZipFile(bad, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")

        with pytest.raises(ValueError, match="xl/workbook.xml"):
            _validate_xlsx(bad)

    def test_docx_validation_rejects_missing_entries(self, tmp_path):
        """_validate_docx rejects a ZIP missing required entries."""
        bad = tmp_path / "bad.docx"
        with zipfile.ZipFile(bad, "w") as zf:
            zf.writestr("random.xml", "<x/>")

        with pytest.raises(ValueError, match="missing required entries"):
            _validate_docx(bad)

    def test_validate_file_nonexistent(self, tmp_path):
        """validate_file raises ValueError for a file that does not exist."""
        missing = tmp_path / "ghost.txt"
        with pytest.raises(ValueError, match="does not exist"):
            validate_file(missing)

    def test_validate_file_unknown_extension_passes(self, tmp_path):
        """Files with unregistered extensions pass validation (no format check)."""
        f = tmp_path / "data.csv"
        f.write_text("a,b,c\n1,2,3")
        # Should not raise -- no CSV validator registered.
        validate_file(f)


class TestComputeSHA256:
    """Standalone SHA-256 utility."""

    def test_compute_sha256_known_value(self, tmp_path):
        """compute_sha256 returns the correct digest for known input."""
        f = tmp_path / "known.bin"
        f.write_bytes(b"hello")
        expected = hashlib.sha256(b"hello").hexdigest()
        assert compute_sha256(f) == expected

    def test_compute_sha256_empty_file(self, tmp_path):
        """compute_sha256 works on an empty file."""
        f = tmp_path / "empty.bin"
        f.write_bytes(b"")
        expected = hashlib.sha256(b"").hexdigest()
        assert compute_sha256(f) == expected


# ═══════════════════════════════════════════════════════════════════════════════
# 2. ErrorStrategy tests
# ═══════════════════════════════════════════════════════════════════════════════

from backend.core.error_strategy import (
    ERROR_RESPONSES,
    DegradationCascade,
    ErrorCategory,
    ErrorResponse,
    classify_error,
)


class TestErrorCategoryClassification:
    """classify_error maps exceptions to the correct category."""

    def test_timeout_error(self):
        assert classify_error(TimeoutError("timed out")) == ErrorCategory.TIMEOUT

    def test_asyncio_timeout(self):
        assert classify_error(asyncio.TimeoutError()) == ErrorCategory.TIMEOUT

    def test_connection_error_is_transient(self):
        assert classify_error(ConnectionError("refused")) == ErrorCategory.TRANSIENT

    def test_memory_error_is_resource(self):
        assert classify_error(MemoryError()) == ErrorCategory.RESOURCE

    def test_file_not_found_is_permanent(self):
        assert classify_error(FileNotFoundError("missing")) == ErrorCategory.PERMANENT

    def test_value_error_is_permanent(self):
        assert classify_error(ValueError("bad input")) == ErrorCategory.PERMANENT

    def test_type_error_is_permanent(self):
        assert classify_error(TypeError("wrong type")) == ErrorCategory.PERMANENT

    def test_unicode_decode_error_is_permanent(self):
        exc = UnicodeDecodeError("utf-8", b"\xff", 0, 1, "invalid")
        assert classify_error(exc) == ErrorCategory.PERMANENT

    def test_unknown_exception_defaults_to_transient(self):
        """An unrecognised exception defaults to TRANSIENT (safe to retry)."""
        assert classify_error(RuntimeError("unknown")) == ErrorCategory.TRANSIENT

    def test_httpx_timeout_by_classname(self):
        """httpx.ReadTimeout matched by class-name string (no import needed)."""
        ReadTimeout = type("ReadTimeout", (Exception,), {"__module__": "httpx._exceptions"})
        exc = ReadTimeout("httpx timeout")
        assert classify_error(exc) == ErrorCategory.TIMEOUT

    def test_httpx_connect_error_is_transient(self):
        """httpx.ConnectError matched by class-name string."""
        ConnectError = type("ConnectError", (Exception,), {"__module__": "httpx._exceptions"})
        exc = ConnectError("connection failed")
        assert classify_error(exc) == ErrorCategory.TRANSIENT


class TestErrorResponseTable:
    """ERROR_RESPONSES policy table has correct values for each category."""

    def test_all_categories_have_responses(self):
        """Every ErrorCategory value has a corresponding ErrorResponse."""
        for cat in ErrorCategory:
            assert cat in ERROR_RESPONSES, f"Missing response for {cat}"

    def test_transient_retries(self):
        r = ERROR_RESPONSES[ErrorCategory.TRANSIENT]
        assert r.should_retry is True
        assert r.max_retries == 3

    def test_permanent_no_retry(self):
        r = ERROR_RESPONSES[ErrorCategory.PERMANENT]
        assert r.should_retry is False
        assert r.max_retries == 0
        assert r.fallback_strategy is None

    def test_injection_quarantines(self):
        r = ERROR_RESPONSES[ErrorCategory.INJECTION_DETECTED]
        assert r.should_retry is False
        assert r.fallback_strategy == "halt_quarantine"

    def test_timeout_simplifies(self):
        r = ERROR_RESPONSES[ErrorCategory.TIMEOUT]
        assert r.fallback_strategy == "simplify_plan"
        assert r.max_retries == 1

    def test_resource_falls_back_to_model(self):
        r = ERROR_RESPONSES[ErrorCategory.RESOURCE]
        assert r.fallback_strategy == "fallback_model"

    def test_user_notification_has_placeholders(self):
        """All user notifications contain at least {job_id} and {node_id}."""
        for cat, resp in ERROR_RESPONSES.items():
            assert "{job_id}" in resp.user_notification, f"{cat} missing {{job_id}}"
            assert "{node_id}" in resp.user_notification, f"{cat} missing {{node_id}}"


class TestDegradationCascade:
    """DegradationCascade resolves correct actions for category + attempt."""

    def setup_method(self):
        self.cascade = DegradationCascade()

    def test_transient_first_attempt_retries(self):
        assert self.cascade.next_action(ErrorCategory.TRANSIENT, 0) == "retry"

    def test_transient_exhausted_dead_letters(self):
        assert self.cascade.next_action(ErrorCategory.TRANSIENT, 3) == "dead_letter"

    def test_transient_past_end_dead_letters(self):
        """Attempts past cascade length always dead-letter."""
        assert self.cascade.next_action(ErrorCategory.TRANSIENT, 100) == "dead_letter"

    def test_permanent_immediately_dead_letters(self):
        assert self.cascade.next_action(ErrorCategory.PERMANENT, 0) == "dead_letter"

    def test_timeout_retries_then_simplifies(self):
        assert self.cascade.next_action(ErrorCategory.TIMEOUT, 0) == "retry"
        assert self.cascade.next_action(ErrorCategory.TIMEOUT, 1) == "simplify"

    def test_model_falls_back(self):
        assert self.cascade.next_action(ErrorCategory.MODEL, 0) == "retry"
        assert self.cascade.next_action(ErrorCategory.MODEL, 2) == "fallback_model"

    def test_contract_partial_delivery(self):
        assert self.cascade.next_action(ErrorCategory.CONTRACT, 1) == "partial_delivery"

    def test_get_response_returns_correct_type(self):
        resp = self.cascade.get_response(ErrorCategory.TRANSIENT)
        assert isinstance(resp, ErrorResponse)
        assert resp.category == ErrorCategory.TRANSIENT

    def test_backoff_for_transient(self):
        assert self.cascade.backoff_for(ErrorCategory.TRANSIENT, 0) == 1_000
        assert self.cascade.backoff_for(ErrorCategory.TRANSIENT, 1) == 5_000
        assert self.cascade.backoff_for(ErrorCategory.TRANSIENT, 2) == 15_000

    def test_backoff_past_list_repeats_last(self):
        """Attempts past the backoff list length repeat the last value."""
        assert self.cascade.backoff_for(ErrorCategory.TRANSIENT, 99) == 15_000

    def test_backoff_empty_returns_zero(self):
        """Categories with no backoff_ms return 0."""
        assert self.cascade.backoff_for(ErrorCategory.PERMANENT, 0) == 0

    def test_injection_dead_letters_immediately(self):
        assert self.cascade.next_action(ErrorCategory.INJECTION_DETECTED, 0) == "dead_letter"

    def test_upstream_dead_letters_immediately(self):
        assert self.cascade.next_action(ErrorCategory.UPSTREAM, 0) == "dead_letter"


# ═══════════════════════════════════════════════════════════════════════════════
# 3. PptxTool tests
# ═══════════════════════════════════════════════════════════════════════════════


def _passthrough_safe_path(file_path: str) -> Path:
    """Bypass safe_resolve jail so tmp_path files work in tests."""
    return Path(file_path).resolve()


def _make_pptx(path: Path, slide_count: int = 1):
    """Create a minimal PPTX file using python-pptx (importskip guarded)."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[0]  # Title Slide
    for i in range(slide_count):
        slide = prs.slides.add_slide(layout)
        title_shape = slide.placeholders[0]
        title_shape.text = f"Slide {i + 1} Title"
        if 1 in slide.placeholders:
            slide.placeholders[1].text = f"Subtitle for slide {i + 1}"
    prs.save(str(path))
    return path


@patch("backend.tools.pptx_tool._safe_path", side_effect=_passthrough_safe_path)
class TestPptxToolSync:
    """Test the synchronous worker functions of pptx_tool directly."""

    def test_read_metadata_success(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_read_metadata

        pptx_file = _make_pptx(tmp_path / "deck.pptx", slide_count=3)
        result = _do_read_metadata(str(pptx_file))

        assert result["success"] is True
        assert result["result"]["slide_count"] == 3
        assert "dimensions" in result["result"]
        assert result["result"]["dimensions"]["width_emu"] is not None

    def test_read_metadata_file_not_found(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_read_metadata

        result = _do_read_metadata(str(tmp_path / "nonexistent.pptx"))
        assert result["success"] is False
        assert "not found" in result["error"].lower()

    def test_extract_content_all_slides(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_extract_content

        pptx_file = _make_pptx(tmp_path / "content.pptx", slide_count=2)
        result = _do_extract_content(str(pptx_file), slide_indices=None)

        assert result["success"] is True
        assert len(result["result"]) == 2
        assert result["result"][0]["slide_index"] == 0

    def test_extract_content_specific_slide(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_extract_content

        pptx_file = _make_pptx(tmp_path / "content2.pptx", slide_count=3)
        result = _do_extract_content(str(pptx_file), slide_indices=[1])

        assert result["success"] is True
        assert len(result["result"]) == 1
        assert result["result"][0]["slide_index"] == 1

    def test_extract_content_invalid_index(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_extract_content

        pptx_file = _make_pptx(tmp_path / "content3.pptx", slide_count=1)
        result = _do_extract_content(str(pptx_file), slide_indices=[5])

        assert result["success"] is False
        assert "out of range" in result["error"]

    def test_extract_styles(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_extract_styles

        pptx_file = _make_pptx(tmp_path / "styles.pptx", slide_count=1)
        result = _do_extract_styles(str(pptx_file))

        assert result["success"] is True
        assert "layouts" in result["result"]
        assert isinstance(result["result"]["layouts"], list)
        assert "style_notes" in result["result"]

    def test_edit_slide_changes_text(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from pptx import Presentation

        from backend.tools.pptx_tool import _do_edit_slide

        pptx_file = _make_pptx(tmp_path / "edit.pptx", slide_count=1)

        # Find the title shape name for editing
        prs = Presentation(str(pptx_file))
        slide = prs.slides[0]
        title_name = slide.placeholders[0].name

        out = tmp_path / "edited.pptx"
        result = _do_edit_slide(
            str(pptx_file),
            slide_index=0,
            edits=[{"shape_name": title_name, "new_text": "Updated Title"}],
            output_path=str(out),
        )

        assert result["success"] is True
        assert result["result"]["sha256"]
        assert result["result"]["size_bytes"] > 0

        # Verify the text was changed
        prs2 = Presentation(str(out))
        assert prs2.slides[0].placeholders[0].text == "Updated Title"

    def test_edit_slide_invalid_index(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import _do_edit_slide

        pptx_file = _make_pptx(tmp_path / "edit2.pptx", slide_count=1)
        result = _do_edit_slide(
            str(pptx_file),
            slide_index=99,
            edits=[{"shape_name": "X", "new_text": "Y"}],
            output_path=None,
        )
        assert result["success"] is False
        assert "out of range" in result["error"]

    def test_apply_edits_multi_slide(self, _mock_sp, tmp_path):
        pytest.importorskip("pptx")
        from pptx import Presentation

        from backend.tools.pptx_tool import _do_apply_edits

        pptx_file = _make_pptx(tmp_path / "multi.pptx", slide_count=2)

        prs = Presentation(str(pptx_file))
        name0 = prs.slides[0].placeholders[0].name
        name1 = prs.slides[1].placeholders[0].name

        out = tmp_path / "multi_edited.pptx"
        result = _do_apply_edits(
            str(pptx_file),
            edit_plan=[
                {"slide_index": 0, "edits": [{"shape_name": name0, "new_text": "A"}]},
                {"slide_index": 1, "edits": [{"shape_name": name1, "new_text": "B"}]},
            ],
            output_path=str(out),
        )

        assert result["success"] is True
        assert 0 in result["result"]["slides_edited"]
        assert 1 in result["result"]["slides_edited"]


class TestPptxToolPathSafety:
    """Path jailing in pptx_tool._safe_path."""

    def test_safe_path_resolves_absolute(self, tmp_path):
        """_safe_path resolves to an absolute path."""
        from backend.tools.pptx_tool import _safe_path

        f = tmp_path / "test.pptx"
        f.write_bytes(b"")
        resolved = _safe_path(str(f))
        assert resolved.is_absolute()

    def test_safe_path_with_mocked_security(self, tmp_path):
        """When safe_resolve is available, it is called with WORKSPACE_ROOT."""
        from backend.tools import pptx_tool

        sentinel = tmp_path / "jailed.pptx"

        mock_safe = MagicMock(return_value=sentinel)
        with patch.dict("sys.modules", {
            "backend.config": MagicMock(WORKSPACE_ROOT=tmp_path),
            "backend.security.paths": MagicMock(safe_resolve=mock_safe, SecurityError=Exception),
        }):
            pptx_tool._safe_path("jailed.pptx")
            mock_safe.assert_called_once()


class TestPptxToolClass:
    """Test the PptxTool class async execute method."""

    def test_tool_name_and_description(self):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import PptxTool

        tool = PptxTool()
        assert tool.name == "pptx"
        assert "PowerPoint" in tool.description

    def test_execute_unknown_action(self):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import PptxTool

        tool = PptxTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="nonexistent", file_path="x.pptx")
        )
        assert result["success"] is False
        assert "Unknown action" in result["error"]

    def test_execute_missing_file_path(self):
        pytest.importorskip("pptx")
        from backend.tools.pptx_tool import PptxTool

        tool = PptxTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="read_metadata", file_path="")
        )
        assert result["success"] is False


# ═══════════════════════════════════════════════════════════════════════════════
# 4. ExcelTool tests
# ═══════════════════════════════════════════════════════════════════════════════


def _make_xlsx(path: Path, data=None):
    """Create a minimal XLSX file using openpyxl (importskip guarded)."""
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    if data is None:
        data = [["Name", "Age"], ["Alice", 30], ["Bob", 25]]
    for row in data:
        ws.append(row)
    wb.save(str(path))
    return path


@patch("backend.tools.excel_tool._safe_path", side_effect=_passthrough_safe_path)
class TestExcelToolSync:
    """Test the synchronous worker functions of excel_tool directly."""

    def test_read_sheet(self, _mock_sp, tmp_path):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _do_read_sheet

        xlsx = _make_xlsx(tmp_path / "data.xlsx")
        result = _do_read_sheet(str(xlsx), sheet_name=None)

        assert result["success"] is True
        assert result["result"]["sheet_name"] == "Sheet1"
        assert result["result"]["headers"] == ["Name", "Age"]
        assert len(result["result"]["rows"]) == 2

    def test_read_range(self, _mock_sp, tmp_path):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _do_read_range

        xlsx = _make_xlsx(tmp_path / "range.xlsx")
        result = _do_read_range(str(xlsx), "A1:B2", sheet_name=None)

        assert result["success"] is True
        assert result["result"]["range"] == "A1:B2"
        assert len(result["result"]["values"]) == 2

    def test_list_sheets(self, _mock_sp, tmp_path):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _do_list_sheets

        xlsx = _make_xlsx(tmp_path / "sheets.xlsx")
        result = _do_list_sheets(str(xlsx))

        assert result["success"] is True
        assert "Sheet1" in result["result"]["sheets"]
        assert result["result"]["sheet_count"] == 1

    def test_get_metadata(self, _mock_sp, tmp_path):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _do_get_metadata

        xlsx = _make_xlsx(tmp_path / "meta.xlsx")
        result = _do_get_metadata(str(xlsx))

        assert result["success"] is True
        assert result["result"]["sheet_count"] == 1
        assert "author" in result["result"]

    def test_extract_formulas(self, _mock_sp, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl

        from backend.tools.excel_tool import _do_extract_formulas

        xlsx_path = tmp_path / "formulas.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = "=SUM(A1:A2)"
        wb.save(str(xlsx_path))

        result = _do_extract_formulas(str(xlsx_path), sheet_name=None)

        assert result["success"] is True
        assert result["result"]["formula_count"] == 1
        assert result["result"]["formulas"][0]["formula"] == "=SUM(A1:A2)"

    def test_file_not_found(self, _mock_sp, tmp_path):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _do_list_sheets

        with pytest.raises(FileNotFoundError):
            _do_list_sheets(str(tmp_path / "missing.xlsx"))

    def test_parse_range_invalid(self, _mock_sp):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _parse_range

        with pytest.raises(ValueError, match="Invalid range"):
            _parse_range("not-a-range")

    def test_parse_range_valid(self, _mock_sp):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import _parse_range

        start, end = _parse_range("A1:D10")
        assert start == "A1"
        assert end == "D10"


class TestExcelToolPathSafety:
    """Path jailing in excel_tool._safe_path."""

    def test_safe_path_resolves_absolute(self, tmp_path):
        from backend.tools.excel_tool import _safe_path

        f = tmp_path / "test.xlsx"
        f.write_bytes(b"")
        resolved = _safe_path(str(f))
        assert resolved.is_absolute()


class TestExcelToolClass:
    """Test the ExcelTool class async execute method."""

    def test_tool_name(self):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import ExcelTool

        tool = ExcelTool()
        assert tool.name == "excel"

    def test_unknown_action(self):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import ExcelTool

        tool = ExcelTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="nonexistent", file_path="x.xlsx")
        )
        assert result["success"] is False
        assert "Unknown action" in result["error"]

    def test_missing_action(self):
        pytest.importorskip("openpyxl")
        from backend.tools.excel_tool import ExcelTool

        tool = ExcelTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="", file_path="x.xlsx")
        )
        assert result["success"] is False


# ═══════════════════════════════════════════════════════════════════════════════
# 5. WordTool tests
# ═══════════════════════════════════════════════════════════════════════════════


def _make_docx(path: Path, paragraphs=None):
    """Create a minimal DOCX file using python-docx (importskip guarded)."""
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    if paragraphs is None:
        paragraphs = ["First paragraph.", "Second paragraph.", "Third paragraph."]
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


@patch("backend.tools.word_tool._safe_path", side_effect=_passthrough_safe_path)
class TestWordToolSync:
    """Test the synchronous worker functions of word_tool directly."""

    def test_read_document(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_read_document

        docx_file = _make_docx(tmp_path / "doc.docx")
        result = _do_read_document(str(docx_file))

        assert result["success"] is True
        assert result["result"]["paragraph_count"] >= 3
        texts = [p["text"] for p in result["result"]["paragraphs"]]
        assert "First paragraph." in texts

    def test_get_paragraph(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_get_paragraph

        docx_file = _make_docx(tmp_path / "para.docx")
        result = _do_get_paragraph(str(docx_file), index=1)

        assert result["success"] is True
        assert result["result"]["text"] == "Second paragraph."
        assert result["result"]["index"] == 1

    def test_get_paragraph_out_of_range(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_get_paragraph

        docx_file = _make_docx(tmp_path / "para2.docx")
        result = _do_get_paragraph(str(docx_file), index=999)

        assert result["success"] is False
        assert "out of range" in result["error"]

    def test_extract_styles(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_extract_styles

        docx_file = _make_docx(tmp_path / "styles.docx")
        result = _do_extract_styles(str(docx_file))

        assert result["success"] is True
        assert isinstance(result["result"]["styles"], list)
        assert result["result"]["total_styles"] > 0

    def test_get_tables_empty(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_get_tables

        docx_file = _make_docx(tmp_path / "no_tables.docx")
        result = _do_get_tables(str(docx_file))

        assert result["success"] is True
        assert result["result"]["table_count"] == 0

    def test_get_tables_with_table(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from docx import Document

        from backend.tools.word_tool import _do_get_tables

        docx_path = tmp_path / "with_table.docx"
        doc = Document()
        doc.add_paragraph("Before table.")
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Header1"
        table.cell(0, 1).text = "Header2"
        table.cell(0, 2).text = "Header3"
        table.cell(1, 0).text = "Val1"
        table.cell(1, 1).text = "Val2"
        table.cell(1, 2).text = "Val3"
        doc.save(str(docx_path))

        result = _do_get_tables(str(docx_path))
        assert result["success"] is True
        assert result["result"]["table_count"] == 1
        assert result["result"]["tables"][0]["row_count"] == 2

    def test_get_metadata(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_get_metadata

        docx_file = _make_docx(tmp_path / "meta.docx")
        result = _do_get_metadata(str(docx_file))

        assert result["success"] is True
        assert "word_count" in result["result"]
        assert "paragraph_count" in result["result"]
        assert result["result"]["paragraph_count"] >= 3

    def test_file_not_found(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _do_read_document

        with pytest.raises(FileNotFoundError):
            _do_read_document(str(tmp_path / "ghost.docx"))

    def test_not_docx_extension(self, _mock_sp, tmp_path):
        pytest.importorskip("docx")
        from backend.tools.word_tool import _open_document

        txt_file = tmp_path / "readme.txt"
        txt_file.write_text("hello")

        with pytest.raises(ValueError, match="Not a .docx file"):
            _open_document(str(txt_file))


class TestWordToolPathSafety:
    """Path jailing in word_tool._safe_path."""

    def test_safe_path_resolves_absolute(self, tmp_path):
        from backend.tools.word_tool import _safe_path

        f = tmp_path / "test.docx"
        f.write_bytes(b"")
        resolved = _safe_path(str(f))
        assert resolved.is_absolute()


class TestWordToolClass:
    """Test the WordTool class async execute method."""

    def test_tool_name(self):
        pytest.importorskip("docx")
        from backend.tools.word_tool import WordTool

        tool = WordTool()
        assert tool.name == "word"

    def test_unknown_action(self):
        pytest.importorskip("docx")
        from backend.tools.word_tool import WordTool

        tool = WordTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="nonexistent", file_path="x.docx")
        )
        assert result["success"] is False
        assert "Unknown action" in result["error"]

    def test_missing_file_path(self):
        pytest.importorskip("docx")
        from backend.tools.word_tool import WordTool

        tool = WordTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="read_document", file_path="")
        )
        assert result["success"] is False


# ═══════════════════════════════════════════════════════════════════════════════
# 6. PDFTool tests
# ═══════════════════════════════════════════════════════════════════════════════


def _make_pdf(path: Path, pages: int = 1, text: str = "Hello from page"):
    """Create a minimal PDF file using pymupdf (importskip guarded)."""
    fitz = pytest.importorskip("fitz")
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"{text} {i + 1}")
    doc.save(str(path))
    doc.close()
    return path


@patch("backend.tools.pdf_tool._safe_path", side_effect=_passthrough_safe_path)
class TestPdfToolSync:
    """Test the synchronous worker functions of pdf_tool directly."""

    def test_extract_text_all_pages(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_extract_text

        pdf = _make_pdf(tmp_path / "text.pdf", pages=3)
        result = _do_extract_text(str(pdf))

        assert result["success"] is True
        assert result["result"]["page_count"] == 3
        assert len(result["result"]["pages"]) == 3
        assert "Hello from page 1" in result["result"]["pages"][0]["text"]

    def test_extract_text_specific_pages(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_extract_text

        pdf = _make_pdf(tmp_path / "text2.pdf", pages=5)
        result = _do_extract_text(str(pdf), pages=[0, 2])

        assert result["success"] is True
        assert len(result["result"]["pages"]) == 2
        assert result["result"]["pages"][0]["page_num"] == 0
        assert result["result"]["pages"][1]["page_num"] == 2

    def test_extract_text_invalid_page(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_extract_text

        pdf = _make_pdf(tmp_path / "text3.pdf", pages=1)
        with pytest.raises(ValueError, match="out of range"):
            _do_extract_text(str(pdf), pages=[99])

    def test_page_count(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_page_count

        pdf = _make_pdf(tmp_path / "count.pdf", pages=7)
        result = _do_page_count(str(pdf))

        assert result["success"] is True
        assert result["result"]["page_count"] == 7

    def test_get_page_text(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_get_page_text

        pdf = _make_pdf(tmp_path / "single.pdf", pages=3, text="TestContent")
        result = _do_get_page_text(str(pdf), page_num=1)

        assert result["success"] is True
        assert result["result"]["page_num"] == 1
        assert "TestContent 2" in result["result"]["text"]

    def test_get_metadata(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_get_metadata

        pdf = _make_pdf(tmp_path / "meta.pdf", pages=2)
        result = _do_get_metadata(str(pdf))

        assert result["success"] is True
        assert result["result"]["page_count"] == 2
        assert "file_size" in result["result"]
        assert result["result"]["file_size"] > 0

    def test_get_toc(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_get_toc

        pdf = _make_pdf(tmp_path / "toc.pdf", pages=1)
        result = _do_get_toc(str(pdf))

        assert result["success"] is True
        assert isinstance(result["result"]["entries"], list)

    def test_search_text_found(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_search_text

        pdf = _make_pdf(tmp_path / "search.pdf", pages=2, text="UniqueMarker")
        result = _do_search_text(str(pdf), "UniqueMarker")

        assert result["success"] is True
        assert result["result"]["match_count"] >= 1

    def test_file_not_found(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_extract_text

        with pytest.raises(FileNotFoundError):
            _do_extract_text(str(tmp_path / "nonexistent.pdf"))

    def test_not_a_pdf(self, _mock_sp, tmp_path):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import _do_extract_text

        txt = tmp_path / "fake.pdf"
        txt.write_text("This is not a PDF")

        with pytest.raises(ValueError):
            _do_extract_text(str(txt))


class TestPdfToolPathSafety:
    """Path jailing in pdf_tool._safe_path."""

    def test_safe_path_resolves_absolute(self, tmp_path):
        from backend.tools.pdf_tool import _safe_path

        f = tmp_path / "test.pdf"
        f.write_bytes(b"")
        resolved = _safe_path(str(f))
        assert resolved.is_absolute()


class TestPdfToolClass:
    """Test the PDFTool class async execute method."""

    def test_tool_name(self):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import PDFTool

        tool = PDFTool()
        assert tool.name == "pdf"

    def test_unknown_action(self):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import PDFTool

        tool = PDFTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="nonexistent", file_path="x.pdf")
        )
        assert result["success"] is False
        assert "Unknown action" in result["error"]

    def test_missing_file_path(self):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import PDFTool

        tool = PDFTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="extract_text", file_path="")
        )
        assert result["success"] is False

    def test_page_text_requires_page_num(self):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import PDFTool

        tool = PDFTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="get_page_text", file_path="x.pdf")
        )
        assert result["success"] is False
        assert "page_num" in result["error"]

    def test_search_text_requires_query(self):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import PDFTool

        tool = PDFTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="search_text", file_path="x.pdf")
        )
        assert result["success"] is False
        assert "query" in result["error"]

    def test_extract_images_requires_output_dir(self):
        pytest.importorskip("fitz")
        from backend.tools.pdf_tool import PDFTool

        tool = PDFTool()
        result = asyncio.get_event_loop().run_until_complete(
            tool.execute(action="extract_images", file_path="x.pdf")
        )
        assert result["success"] is False
        assert "output_dir" in result["error"]
