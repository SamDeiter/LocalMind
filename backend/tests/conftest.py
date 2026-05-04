"""
Shared fixtures for LocalMind document-tool integration tests.

Generates real PPTX / XLSX / DOCX / PDF sample files programmatically so that
the tests exercise the actual file-format libraries (python-pptx, openpyxl,
python-docx, pymupdf) rather than mocks.
"""

import os
import shutil
import sys
from pathlib import Path

import pytest

# ---------------------------------------------------------------------------
# Ensure project root is on sys.path so `backend.*` imports work.
# ---------------------------------------------------------------------------
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


# ---------------------------------------------------------------------------
# Workspace fixture — every test gets its own temp directory
# ---------------------------------------------------------------------------


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    """Return a fresh temporary directory that acts as a sandboxed workspace."""
    ws = tmp_path / "workspace"
    ws.mkdir()
    return ws


# ---------------------------------------------------------------------------
# Bypass the security jail in tool _safe_path helpers.
#
# The document tools try to import backend.config.WORKSPACE_ROOT and
# backend.security.paths.safe_resolve, which jails all paths to
# ~/LocalMind_Workspace.  In tests we use tmp_path, so we monkeypatch
# _safe_path to just resolve the path directly.
# ---------------------------------------------------------------------------

def _passthrough_safe_path(file_path: str) -> Path:
    """Test-only replacement for _safe_path — resolves without jailing."""
    return Path(file_path).resolve()


@pytest.fixture(autouse=True)
def _bypass_safe_path(monkeypatch):
    """Auto-applied fixture that patches _safe_path in every tool module."""
    import backend.tools.pptx_tool as pptx_mod
    import backend.tools.excel_tool as excel_mod
    import backend.tools.word_tool as word_mod
    import backend.tools.pdf_tool as pdf_mod

    monkeypatch.setattr(pptx_mod, "_safe_path", _passthrough_safe_path)
    monkeypatch.setattr(excel_mod, "_safe_path", _passthrough_safe_path)
    monkeypatch.setattr(word_mod, "_safe_path", _passthrough_safe_path)
    monkeypatch.setattr(pdf_mod, "_safe_path", _passthrough_safe_path)


# ---------------------------------------------------------------------------
# PPTX fixture
# ---------------------------------------------------------------------------


def _create_sample_pptx(dest: Path) -> Path:
    """Create a 3-slide PowerPoint with titled text and some formatting."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()

    # Slide 1 — Title Slide
    layout_title = prs.slide_layouts[0]  # Title Slide layout
    slide1 = prs.slides.add_slide(layout_title)
    slide1.shapes.title.text = "Integration Test Deck"
    subtitle = slide1.placeholders[1]
    subtitle.text = "Auto-generated for testing"

    # Slide 2 — Title and Content layout
    layout_content = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(layout_content)
    slide2.shapes.title.text = "Slide Two Title"
    body = slide2.placeholders[1]
    tf = body.text_frame
    tf.text = "Bullet point one"
    p2 = tf.add_paragraph()
    p2.text = "Bullet point two"

    # Slide 3 — Blank with a text box that has explicit font formatting
    layout_blank = prs.slide_layouts[6]  # Blank
    slide3 = prs.slides.add_slide(layout_blank)
    from pptx.util import Emu
    txBox = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tf3 = txBox.text_frame
    p = tf3.paragraphs[0]
    run = p.add_run()
    run.text = "Formatted text"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    run.font.name = "Arial"

    prs.save(str(dest))
    return dest


@pytest.fixture()
def sample_pptx(workspace: Path) -> Path:
    """Return the path to a freshly-created 3-slide PPTX in the workspace."""
    return _create_sample_pptx(workspace / "sample.pptx")


# ---------------------------------------------------------------------------
# XLSX fixture
# ---------------------------------------------------------------------------


def _create_sample_xlsx(dest: Path) -> Path:
    """Create an Excel workbook with 2 sheets and some data."""
    import openpyxl

    wb = openpyxl.Workbook()

    # Sheet 1 (default) — "Sales"
    ws1 = wb.active
    ws1.title = "Sales"
    ws1.append(["Product", "Q1", "Q2", "Q3", "Q4"])
    ws1.append(["Widget A", 100, 150, 200, 250])
    ws1.append(["Widget B", 80, 120, 160, 200])
    ws1.append(["Widget C", 60, 90, 130, 170])
    # Add a formula
    ws1["F1"] = "Total"
    ws1["F2"] = "=SUM(B2:E2)"

    # Sheet 2 — "Metadata"
    ws2 = wb.create_sheet(title="Metadata")
    ws2.append(["Key", "Value"])
    ws2.append(["Author", "Test Suite"])
    ws2.append(["Version", "1.0"])

    wb.save(str(dest))
    return dest


@pytest.fixture()
def sample_xlsx(workspace: Path) -> Path:
    """Return the path to a freshly-created XLSX in the workspace."""
    return _create_sample_xlsx(workspace / "sample.xlsx")


# ---------------------------------------------------------------------------
# DOCX fixture
# ---------------------------------------------------------------------------


def _create_sample_docx(dest: Path) -> Path:
    """Create a Word document with paragraphs, formatting, and a table."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.core_properties.author = "Test Suite"
    doc.core_properties.title = "Integration Test Document"

    # Heading
    doc.add_heading("Document Title", level=0)

    # Normal paragraphs
    p1 = doc.add_paragraph("This is the first paragraph of the test document.")
    p2 = doc.add_paragraph("")
    run = p2.add_run("Bold and formatted text")
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Calibri"

    doc.add_paragraph("Third paragraph for content extraction tests.")

    # A small table
    table = doc.add_table(rows=3, cols=3)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"R{i}C{j}"

    doc.save(str(dest))
    return dest


@pytest.fixture()
def sample_docx(workspace: Path) -> Path:
    """Return the path to a freshly-created DOCX in the workspace."""
    return _create_sample_docx(workspace / "sample.docx")


# ---------------------------------------------------------------------------
# PDF fixture
# ---------------------------------------------------------------------------


def _create_sample_pdf(dest: Path) -> Path:
    """Create a simple 2-page PDF using pymupdf (fitz)."""
    import fitz

    doc = fitz.open()

    # Page 1
    page1 = doc.new_page(width=612, height=792)
    page1.insert_text((72, 72), "PDF Integration Test — Page 1", fontsize=16)
    page1.insert_text((72, 120), "This is sample content for testing the PDF tool.")

    # Page 2
    page2 = doc.new_page(width=612, height=792)
    page2.insert_text((72, 72), "PDF Integration Test — Page 2", fontsize=16)
    page2.insert_text((72, 120), "Second page with more content for search tests.")

    doc.save(str(dest))
    doc.close()
    return dest


@pytest.fixture()
def sample_pdf(workspace: Path) -> Path:
    """Return the path to a freshly-created 2-page PDF in the workspace."""
    return _create_sample_pdf(workspace / "sample.pdf")


# ---------------------------------------------------------------------------
# Convenience: all sample docs at once
# ---------------------------------------------------------------------------


@pytest.fixture()
def all_samples(workspace: Path) -> dict[str, Path]:
    """Return a dict of all sample document paths keyed by extension."""
    return {
        "pptx": _create_sample_pptx(workspace / "sample.pptx"),
        "xlsx": _create_sample_xlsx(workspace / "sample.xlsx"),
        "docx": _create_sample_docx(workspace / "sample.docx"),
        "pdf": _create_sample_pdf(workspace / "sample.pdf"),
    }
